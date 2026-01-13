#!/usr/bin/env python3
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# 파일 경로 설정
SOURCE_XLSX = "ASCP_Data_Input_201406.xlsx"
DQN_RESULT_XLSX = "output_pairing_201406_1000_exp_gpu0.xlsx"
OUTPUT_XLSX = "input_15656.xlsx"
INITIAL_SOLUTION_XLSX = "output.xlsx"

def pad_sheet(ws, title_row):
    ws.append([title_row])   
    ws.append([""])          
    ws.append([""])          

def build_input():
    wb = Workbook()
    wb.remove(wb.active)

    # 비행 데이터를 가장 먼저 읽도록 (공항 리스트 확보를 위해)
    raw = pd.read_excel(SOURCE_XLSX, sheet_name="User_Flight", header=None)
    raw = raw[raw.iloc[:, 3] != "ORIGIN_DATE"].dropna(subset=[2,3,4,5,6])
    
    # S/N 매핑용 리스트
    sn_list = raw.iloc[:, 0].astype(str).tolist()

    # --- User_Time ---
    ws = wb.create_sheet("User_Time")
    ut = pd.read_excel(SOURCE_XLSX, sheet_name="User_Time", header=None)
    for r in dataframe_to_rows(ut, index=False, header=False):
        ws.append(r)

    # --- Program_Cost ---
    ws = wb.create_sheet("Program_Cost")
    ws.append(["Program Cost"])
    ws.append(["type", "crewNum", "flightCost", "layoverCost", "quickTurnCost"])
    pc = pd.read_excel(SOURCE_XLSX, sheet_name="Program_Cost", header=1)
    pc = pc.iloc[:, :5].loc[pc.iloc[:, 0].notna()]
    for col in [1, 2, 3, 4]:
        pc.iloc[:, col] = pd.to_numeric(pc.iloc[:, col], errors="coerce").fillna(0).astype(int)
    for _, row in pc.iterrows():
        ws.append(row.tolist())

    # --- User_Hotel ---
    ws = wb.create_sheet("User_Hotel")
    pad_sheet(ws, "Airport Hotel Cost")
    uh = pd.read_excel(SOURCE_XLSX, sheet_name="User_Hotel", header=2)
    uh = uh.loc[uh.iloc[:, 0].notna()]
    uh.iloc[:, 0] = uh.iloc[:, 0].astype(str).str.strip()

    # 비행 정보에서 모든 공항 추출 
    flights_airports = set(raw.iloc[:, 2].astype(str).str.strip().unique()) | \
                       set(raw.iloc[:, 4].astype(str).str.strip().unique())
    
    hotel_airports = set(uh.iloc[:, 0].tolist())
    missing_airports = flights_airports - hotel_airports
    existing_airports = flights_airports & hotel_airports 

    print(f"\n[Hotel Data Check]")
    # print(f"   - 원래 인식된 공항 ({len(existing_airports)}개): {sorted(list(existing_airports))}")
    
    if missing_airports:
        print(f"   - 누락되어 추가된 공항 ({len(missing_airports)}개): {sorted(list(missing_airports))}")
        for ap in missing_airports:
            new_row = [ap, 1] + [1] * (len(uh.columns) - 2)
            uh.loc[len(uh)] = new_row
    else:
        print(f"   - 누락된 공항 없음")

    for r in dataframe_to_rows(uh, index=False, header=False):
        ws.append(r)
    
    airport_set = set(uh.iloc[:, 0].astype(str).tolist())

    # --- User_Deadhead ---
    ws = wb.create_sheet("User_Deadhead")
    pad_sheet(ws, "Deadhead Cost")
    ud = pd.read_excel(SOURCE_XLSX, sheet_name="User_Deadhead", header=2)
    ud = ud.loc[ud.iloc[:, 0].astype(str).isin(airport_set) & ud.iloc[:, 1].astype(str).isin(airport_set)]
    for r in dataframe_to_rows(ud, index=False, header=False):
        ws.append(r)

    # --- User_Flight ---
    ws_flight = wb.create_sheet("User_Flight")
    pad_sheet(ws_flight, "Flight Data")
    flight_df = pd.DataFrame({
        "serialNumber": sn_list,
        "tailNumber":   raw.iloc[:, 1].astype(str),
        "origin":       raw.iloc[:, 2],
        "originTime":   pd.to_datetime(raw.iloc[:, 3]),
        "dest":         raw.iloc[:, 4],
        "destTime":     pd.to_datetime(raw.iloc[:, 5]),
        "aircraft":     raw.iloc[:, 6],
    })
    for r in dataframe_to_rows(flight_df, index=False, header=False):
        ws_flight.append(r)

    wb.save(OUTPUT_XLSX)
    print(f"Created {OUTPUT_XLSX}")

# --- Initial_Solution (output.xlsx) ---
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Sheet"
    ws_out.append(["Sheet"]) 

    assigned_flights = set()
    pairing_idx = 1

    try:
        dqn_res = pd.read_excel(DQN_RESULT_XLSX)
        
        for _, row in dqn_res.iterrows():
            converted_row = [pairing_idx] 
            for val in row[1:]:
                if pd.notna(val):
                    flight_id = int(val) 
                    converted_row.append(flight_id)
                    assigned_flights.add(flight_id)
            
            if len(converted_row) > 1:
                ws_out.append(converted_row)
                pairing_idx += 1
        
        all_flight_ids = set(range(len(flight_df)))
        missing_flights = sorted(list(all_flight_ids - assigned_flights))
        
        print(f"📦 누락된 {len(missing_flights)}개의 비행기를 숫자 형식으로 추가")
        for f_id in missing_flights:
            ws_out.append([pairing_idx, f_id])
            pairing_idx += 1
            
        wb_out.save(INITIAL_SOLUTION_XLSX)
        print(f"✅ Created {INITIAL_SOLUTION_XLSX} with Pure Integer IDs.")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    build_input()