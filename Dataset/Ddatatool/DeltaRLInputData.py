import pandas as pd
from openpyxl import load_workbook
from pathlib import Path
from openpyxl.cell.cell import MergedCell
import sys

# =========================
# 상수 설정 (0 금지)
# =========================
DEADHEAD_COST = 50000     # 공항쌍 이동 1회 비용(원)
HOTEL_COST    = 120000    # 공항별 1박 호텔비(원)

# =========================
# 유틸 함수
# =========================
def sn_from(tn: str, origin_dt: pd.Timestamp) -> str:
    return f"{tn}_{origin_dt.strftime('%y%m%d%H%M')}"

def clear_sheet(ws, start_row, start_col, end_col):
    """
    서식/병합은 유지하고 값만 비움.
    병합셀(MergedCell)은 value 수정이 불가하므로 건너뜀.
    """
    for r in range(start_row, ws.max_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None

def safe_get_col(df: pd.DataFrame, candidates):
    for c in candidates:
        if c in df.columns:
            return c
    return None

# =========================
# 메인 변환 함수 (CSV -> ASCP Input)
# =========================
def csv_to_ascp_xlsx(flight_csv: Path, template_xlsx: Path, out_xlsx: Path):
    if not flight_csv.exists():
        print(f"[ERROR] flight csv가 없습니다: {flight_csv}")
        sys.exit(1)

    if not template_xlsx.exists():
        print(f"[ERROR] 템플릿 파일이 없습니다: {template_xlsx}")
        sys.exit(1)

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)

    # 1) flight 로드
    df = pd.read_csv(flight_csv)

    # 2) 필요한 컬럼 찾기 (flight_201901.csv 포맷/다른 포맷 일부 대응)
    origin_airport_col = safe_get_col(df, ["ORIGIN"])
    dest_airport_col   = safe_get_col(df, ["DEST"])
    dep_col            = safe_get_col(df, ["origin", "DEP_TIME"])
    arr_col            = safe_get_col(df, ["dest", "ARR_TIME"])
    aircraft_col       = safe_get_col(df, ["AIRCRAFT_TYPE", "AIRCRAFT_MODEL"])
    tail_col           = safe_get_col(df, ["TAIL_NUM", "Tail_Number", "Tail Number", "TN", "T/N"])
    index_col          = safe_get_col(df, ["INDEX"])

    missing = [("ORIGIN", origin_airport_col), ("DEST", dest_airport_col),
               ("origin/DEP_TIME", dep_col), ("dest/ARR_TIME", arr_col),
               ("AIRCRAFT_TYPE/AIRCRAFT_MODEL", aircraft_col)]
    missing = [name for name, col in missing if col is None]
    if missing:
        print("[ERROR] flight csv에 필요한 컬럼이 부족합니다:", missing)
        print("현재 컬럼:", list(df.columns))
        sys.exit(1)

    # 3) 시간 파싱
    df[dep_col] = pd.to_datetime(df[dep_col], errors="coerce")
    df[arr_col] = pd.to_datetime(df[arr_col], errors="coerce")

    if df[dep_col].isna().any() or df[arr_col].isna().any():
        bad_dep = int(df[dep_col].isna().sum())
        bad_arr = int(df[arr_col].isna().sum())
        print(f"[ERROR] 시간 파싱 실패가 있습니다. dep 실패 {bad_dep}, arr 실패 {bad_arr}")
        sys.exit(1)

    # 4) 공항/기종/쌍 구성
    airports = sorted(set(df[origin_airport_col].astype(str)) | set(df[dest_airport_col].astype(str)))
    aircraft_types = sorted(df[aircraft_col].astype(str).dropna().unique())
    observed_pairs = sorted(set(zip(df[origin_airport_col].astype(str), df[dest_airport_col].astype(str))))

    # 5) 템플릿 로드
    wb = load_workbook(template_xlsx)

    # =========================
    # User_Flight만 채움
    # =========================
    if "User_Flight" not in wb.sheetnames:
        print("[ERROR] 템플릿에 User_Flight 시트가 없습니다.")
        sys.exit(1)

    ws = wb["User_Flight"]
    clear_sheet(ws, start_row=4, start_col=1, end_col=7)

    df_sorted = df.sort_values([dep_col, arr_col]).reset_index(drop=True)

    for i, r0 in enumerate(df_sorted.itertuples(index=False)):
        r = 4 + i
        dep_dt = pd.Timestamp(getattr(r0, dep_col)).to_pydatetime()
        arr_dt = pd.Timestamp(getattr(r0, arr_col)).to_pydatetime()

        if tail_col is not None:
            tn_val = str(getattr(r0, tail_col))
        elif index_col is not None:
            tn_val = str(getattr(r0, index_col))
        else:
            tn_val = f"TN{i+1:06d}"

        origin_air = str(getattr(r0, origin_airport_col))
        dest_air   = str(getattr(r0, dest_airport_col))
        aircraft   = str(getattr(r0, aircraft_col))

        ws.cell(r, 1).value = sn_from(tn_val, pd.Timestamp(dep_dt))  # SN
        ws.cell(r, 2).value = tn_val                                 # T/N
        ws.cell(r, 3).value = origin_air                              # ORIGIN
        ws.cell(r, 4).value = dep_dt                                  # ORIGIN_DATE
        ws.cell(r, 5).value = dest_air                                # DEST
        ws.cell(r, 6).value = arr_dt                                  # DEST_DATE
        ws.cell(r, 7).value = aircraft                                # AIRCRAFT_TYPE

    # =========================
    # User_Deadhead 채움 (상수 비용)
    # =========================
    if "User_Deadhead" not in wb.sheetnames:
        print("[ERROR] 템플릿에 User_Deadhead 시트가 없습니다.")
        sys.exit(1)

    ws = wb["User_Deadhead"]
    clear_sheet(ws, start_row=4, start_col=1, end_col=3)

    rows = []
    for a in airports:
        rows.append((a, a, 0))  # self-loop 0

    for a, b in observed_pairs:
        if a != b:
            rows.append((a, b, DEADHEAD_COST))
            rows.append((b, a, DEADHEAD_COST))

    for i, (a, b, cost) in enumerate(rows):
        r = 4 + i
        ws.cell(r, 1).value = a
        ws.cell(r, 2).value = b
        ws.cell(r, 3).value = cost

    # =========================
    # User_Hotel 채움 (상수 비용)
    # =========================
    if "User_Hotel" not in wb.sheetnames:
        print("[ERROR] 템플릿에 User_Hotel 시트가 없습니다.")
        sys.exit(1)

    ws = wb["User_Hotel"]
    clear_sheet(ws, start_row=4, start_col=1, end_col=2)

    for i, a in enumerate(airports):
        r = 4 + i
        ws.cell(r, 1).value = a
        ws.cell(r, 2).value = HOTEL_COST

    # =========================
    # 저장
    # =========================
    wb.save(out_xlsx)

    print("===================================")
    print(f"[OK] 생성 완료: {out_xlsx}")
    print(f" - Flights  : {len(df_sorted)}")
    print(f" - Airports : {len(airports)}")
    print(f" - Aircraft : {len(aircraft_types)}")
    print(" - NOTE     : User_Time/User_Party/Program_Cost/_Database는 미수정")
    print("===================================")

# =========================
# CLI
# =========================
if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--flight", required=True, help="flight csv 경로")
    p.add_argument("--tpl", required=True, help="ASCP_Data_Input_new.xlsx 템플릿 경로")
    p.add_argument("--out", required=True, help="출력 xlsx 경로")
    args = p.parse_args()

    csv_to_ascp_xlsx(
        flight_csv=Path(args.flight),
        template_xlsx=Path(args.tpl),
        out_xlsx=Path(args.out),
    )